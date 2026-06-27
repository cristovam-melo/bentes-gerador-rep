from docx import Document
import shutil

shutil.copy('template_repo.zip', 'template_repo.docx')
doc = Document('template_repo.docx')

for table in doc.tables:
    if len(table.rows) > 0 and "Destinat" in table.cell(0,0).text:
        print("FOUND DESTINATARIOS TABLE:")
        for r_idx, row in enumerate(table.rows):
            print(f"Row {r_idx}:")
            for c_idx, cell in enumerate(row.cells):
                print(f"  Col {c_idx}: {cell.text.replace(chr(10), ' ')}")
